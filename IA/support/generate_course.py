"""Génère un cours PowerPoint éditable et ses figures, à partir des mesures réelles."""
from pathlib import Path
import json
import os

os.environ.setdefault('MPLCONFIGDIR', '/tmp/cours-matplotlib')
os.environ.setdefault('XDG_CACHE_HOME', '/tmp/cours-cache')
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap
from sklearn.datasets import load_digits
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'support'
OUT = ROOT / 'livrables'
OUT.mkdir(exist_ok=True)
BG = 'F6F7F3'
INK = '153D40'
TEAL = '087F78'
MINT = 'AEF0D6'
ORANGE = 'E79243'
MUTED = '586E72'
LINE = 'DCE6DF'
WHITE = 'FFFFFF'
FONT = 'Arial'
CODE = 'Consolas'
P = Presentation()
P.slide_width, P.slide_height = Inches(13.3333), Inches(7.5)
slides = []


def color(value):
    return RGBColor.from_string(value)


def rect(s, x, y, w, h, fill=WHITE, border=None, radius=True):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                              Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if border:
        shape.line.color.rgb = color(border)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.12
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    return shape


def text(s, x, y, w, h, content, size=20, fg=INK, bold=False, font=FONT,
         align=PP_ALIGN.LEFT):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.01)
    frame.margin_top = frame.margin_bottom = Inches(0.015)
    for index, line in enumerate(str(content).split('\n')):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = align
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color(fg)
        p.line_spacing = 1.12
        p.space_after = Pt(7 if size >= 18 else 3)
    return shape


def line(s, x1, y1, x2, y2, fg=LINE, width=1.2):
    shape = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color(fg)
    shape.line.width = Pt(width)
    return shape


def circle(s, x, y, d, fill=TEAL, border=None):
    shape = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if border:
        shape.line.color.rgb = color(border)
    else:
        shape.line.fill.background()
    return shape


def picture(s, path, x, y, w, h=None):
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    target_h = w * ih / iw
    if h is not None and target_h > h:
        w = h * iw / ih
        target_h = h
    s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(target_h))


def new(title, section, subtitle='', dark=False, notes=''):
    s = P.slides.add_slide(P.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color(INK if dark else BG)
    text(s, .65, .35, 10.5, .3, 'ATELIER PYTHON   /   ' + section.upper(), 10, MINT if dark else TEAL, True)
    text(s, .65, .9, 12, .75, title, 31, WHITE if dark else INK, True)
    if subtitle:
        text(s, .67, 1.72, 11.9, .55, subtitle, 16, 'BCD4CE' if dark else MUTED)
    line(s, .65, 7.06, 12.68, 7.06, '3B5C5C' if dark else LINE)
    text(s, .65, 7.17, 9.8, .2, 'CRÉER UNE IA QUI APPREND  •  NUMPY + PYSIDE6', 8.5, 'BCD4CE' if dark else MUTED)
    slides.append((s, dark))
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def tag(s, x, y, title, w=2.4, dark=False):
    rect(s, x, y, w, .34, '315854' if dark else 'E4F0E9')
    text(s, x+.12, y+.055, w-.2, .23, title, 9, MINT if dark else TEAL, True)


def box(s, x, y, w, h, title, body, accent=TEAL, body_size=18):
    rect(s, x, y, w, h, WHITE, LINE)
    rect(s, x, y+.16, .055, .54, accent, radius=False)
    text(s, x+.23, y+.24, w-.45, .64, title, 21, accent, True)
    if body:
        text(s, x+.23, y+1.04, w-.48, max(.4, h-1.2), body, body_size)


def code(s, x, y, w, h, content, size=16, caption='Extrait du code fourni'):
    rect(s, x, y, w, h, INK)
    text(s, x+.23, y+.15, w-.45, .3, caption.upper(), 9, MINT, True)
    text(s, x+.23, y+.62, w-.46, h-.78, content, size, WHITE, font=CODE)


def takeaway(s, content, y=6.3):
    rect(s, .67, y, 11.98, .53, 'E4F0E9')
    text(s, .86, y+.12, 11.55, .34, content, 15, TEAL, True)


def source(s, caption, url=None):
    shape = text(s, .68, 6.88, 11.8, .16, caption, 7.3, MUTED)
    if url:
        shape.text_frame.paragraphs[0].runs[0].hyperlink.address = url


def neuron_diagram(s, x=7.2, y=2.25, scale=1, stretch=1):
    cols = [(0, 5), (1.9, 4), (3.8, 3)]
    for (left, count), (right, nxt) in zip(cols, cols[1:]):
        for i in range(count):
            for j in range(nxt):
                line(s, x+left*scale*stretch+.23*scale, y+i*(.64 if left == 0 else .78)*scale+.23*scale,
                     x+right*scale*stretch+.23*scale, y+j*.78*scale+.23*scale,
                     '396661', .7)
    for ci, (offset, count) in enumerate(cols):
        for i in range(count):
            circle(s, x+offset*scale*stretch, y+i*(.64 if ci==0 else .78)*scale,
                   .46*scale, MINT if ci!=1 else ORANGE)


def assets():
    digits = load_digits()
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 11,
                         'axes.spines.top': False, 'axes.spines.right': False,
                         'axes.labelcolor': '#'+INK, 'xtick.color':'#'+MUTED,
                         'ytick.color':'#'+MUTED, 'text.color':'#'+INK})
    cmap = LinearSegmentedColormap.from_list('digits', ['#F1F5EE','#087F78'])
    fig, axs = plt.subplots(2, 5, figsize=(9.5, 3.6), facecolor='#'+BG)
    for k, ax in enumerate(axs.flat):
        idx = np.flatnonzero(digits.target==k)[0]
        ax.imshow(digits.images[idx], cmap=cmap, vmin=0, vmax=16, interpolation='nearest')
        ax.set_title(str(k), weight='bold')
        ax.set_axis_off()
    fig.tight_layout()
    fig.savefig(ASSETS/'digits.png', dpi=180, facecolor=fig.get_facecolor())
    plt.close(fig)
    image = digits.images[np.flatnonzero(digits.target==7)[0]]
    fig, ax = plt.subplots(figsize=(4.6, 4.2), facecolor='#'+BG)
    ax.imshow(image, cmap=cmap, vmin=0, vmax=16)
    for r in range(8):
        for c in range(8):
            ax.text(c,r,str(int(image[r,c])),ha='center',va='center',fontsize=12,
                    color='white' if image[r,c]>9 else '#'+INK)
    ax.set_xticks([]); ax.set_yticks([])
    ax.set_title('Un chiffre = une grille de 8 × 8 intensités', fontsize=12)
    fig.tight_layout()
    fig.savefig(ASSETS/'pixels.png', dpi=180, facecolor=fig.get_facecolor())
    plt.close(fig)
    report=json.loads((ASSETS/'metrics.json').read_text())
    h=report['history']; epochs=[v['epoch'] for v in h]
    for metric,title,ylabel in [('loss','L’erreur diminue pendant l’apprentissage','Entropie croisée'),
                                ('accuracy','Suivre les bonnes prédictions','Exactitude (%)')]:
        fig,ax=plt.subplots(figsize=(8,4.2),facecolor='#'+BG)
        ax.set_facecolor('#'+BG)
        for split,col,name in [('train',TEAL,'Apprentissage'),('val',ORANGE,'Validation')]:
            vals=[v[f'{split}_{metric}']*(100 if metric=='accuracy' else 1) for v in h]
            ax.plot(epochs,vals,color='#'+col,label=name,lw=2.5,ls='-' if split=='train' else '--')
        ax.set(xlabel='Époque',ylabel=ylabel,title=title)
        ax.grid(alpha=.2)
        ax.legend(frameon=False,loc='best')
        fig.tight_layout()
        fig.savefig(ASSETS/f'{metric}.png',dpi=180,facecolor=fig.get_facecolor())
        plt.close(fig)
    matrix=np.array(report['confusion_matrix'])
    fig,ax=plt.subplots(figsize=(6.1,5.2),facecolor='#'+BG)
    ax.imshow(matrix,cmap=cmap)
    for r in range(10):
        for c in range(10):
            ax.text(c,r,str(matrix[r,c]),ha='center',va='center',fontsize=10,
                    color='white' if matrix[r,c]>matrix.max()*.55 else '#'+INK)
    ax.set(xticks=range(10),yticks=range(10),xlabel='Chiffre prédit',ylabel='Vrai chiffre')
    fig.tight_layout()
    fig.savefig(ASSETS/'confusion.png',dpi=180,facecolor=fig.get_facecolor())
    plt.close(fig)
    return report


def build(report):
    counts=report['split_counts']; last=report['history'][-1]
    pct=lambda x:f'{100*x:.1f}'.replace('.',',')+' %'
    s=new('Une IA qui apprend.', 'Cours + projet guidé', dark=True,
          notes="Durée conseillée : 2 heures, dont environ 50 minutes de manipulation. Prérequis : savoir lancer un script, lire une fonction et utiliser une boucle. Le cours est un premier réseau de neurones supervisé, pas un chatbot ni un grand modèle de langage. Le code utilise NumPy pour les calculs ; scikit-learn sert seulement à obtenir les données et à les répartir. Présenter l’application finale avant les formules.")
    text(s,.67,2.05,7.1,1.8,'Du neurone\nà l’interface Python.',41,WHITE,True)
    text(s,.72,4.25,6.5,1.0,'Comprendre, entraîner et tester\nun petit réseau de neurones.',22,'BCD4CE')
    tag(s,.72,5.75,'DÉBUTANT +',2.0,True)
    tag(s,2.95,5.75,'ENVIRON 2 HEURES',2.65,True)
    neuron_diagram(s,8.05,2.6,.95)
    text(s,8.05,5.85,4.5,.5,'64 pixels  →  32 neurones  →  10 chiffres',14,MINT)

    s=new('Le projet que vous allez construire', '01 / Le problème',
          'Une application qui apprend à classer les chiffres manuscrits de 0 à 9.',
          notes="Démonstration, 3 min. Montrer le bouton d’entraînement, la progression puis une prédiction. Avant l’apprentissage, les poids sont tirés au hasard. Après les mises à jour, le réseau fait mieux sur des images qu’il n’a pas apprises. L’application livrée est une petite application de bureau : elle se lance sur un ordinateur avec un affichage graphique, pas dans un navigateur.")
    picture(s,ASSETS/'interface.png',.68,2.45,8.0,4.1)
    box(s,9.0,2.45,3.65,1.08,'01  Apprendre','',TEAL)
    text(s,9.22,3.6,3.12,.7,'Ajuster les poids\nà partir d’exemples.',17)
    text(s,9.22,4.5,3.12,.55,'02  Prédire',21,TEAL,True)
    text(s,9.22,5.16,3.12,1.0,'Recevoir une image et\nproposer un chiffre.',17)

    s=new('Votre parcours en quatre étapes', '01 / Le problème',
          'Vous n’avez pas besoin de maîtriser le calcul différentiel pour commencer.',
          notes="3 min. Demander ce qu’est une variable et une fonction. Présenter les quatre étapes et rappeler les pauses d’exercice. L’algèbre matricielle est introduite visuellement ; la rétropropagation détaillée est une lecture facultative. Pour un groupe novice, passer plus de temps sur les données et l’entraînement que sur la dérivation des gradients.")
    for i,(t,b) in enumerate([('Comprendre','Données, neurones\net poids.\n≈ 25 min'),('Programmer','Propagation, erreur\net mise à jour.\n≈ 35 min'),('Observer','Courbes, validation\net petit écran Qt.\n≈ 25 min'),('Expérimenter','Exercices, erreurs\net bilan final.\n≈ 35 min')]):
        box(s,.68+i*3.06,2.6,2.83,3.5,f'{i+1:02d}  {t}',b,TEAL if i%2==0 else ORANGE,18)
    takeaway(s,'Objectif : savoir expliquer ce qui change quand le modèle « apprend ».')

    s=new('Apprendre = ajuster des nombres', '01 / Le problème',
          'Le programme connaît la méthode. Les exemples déterminent les poids.',
          notes="4 min. Un exemple est un couple entrée x, étiquette y. Ici x contient les pixels et y le chiffre correct. Le programme n’écrit pas dix règles à la main : il cherche les poids qui réduisent une erreur. Distinguer apprentissage et inférence : prédire seul n’ajuste aucun poids. Il faut relancer un entraînement pour apprendre de nouveaux exemples.")
    box(s,.68,2.65,3.4,3.15,'Entrée x','Une image de chiffre.\n\n64 valeurs numériques.')
    text(s,4.23,3.63,.6,.7,'→',34,TEAL,True)
    box(s,4.9,2.65,3.4,3.15,'Modèle fθ','Des poids θ.\n\nUne transformation calculée.')
    text(s,8.45,3.63,.6,.7,'→',34,TEAL,True)
    box(s,9.13,2.65,3.5,3.15,'Sortie ŷ','10 scores normalisés.\n\nLe chiffre le plus probable.')
    takeaway(s,'L’étiquette y permet de mesurer l’erreur entre la réponse attendue et la prédiction.')

    s=new('Notre matière première : 1 797 images', '01 / Le problème',
          'load_digits() fournit un petit jeu de données directement avec scikit-learn.',
          notes="4 min. Les exemples sont des chiffres écrits à la main, en niveaux de gris. Ce jeu de données n’est pas MNIST : les images font 8×8, pas 28×28. Les intensités sont des entiers entre 0 et 16. Il est disponible sans téléchargement supplémentaire après installation de scikit-learn. Les données ne garantissent pas de bonnes prédictions sur des photos prises au téléphone.")
    picture(s,ASSETS/'digits.png',.65,2.5,8.55,3.7)
    box(s,9.45,2.65,3.2,3.25,'À retenir','8 × 8 pixels\n\n10 classes : 0 à 9\n\nValeurs : 0 à 16',body_size=19)
    source(s,'Source : scikit-learn, load_digits — jeu de chiffres manuscrits 8×8.',
           'https://scikit-learn.org/stable/modules/generated/sklearn.datasets.load_digits.html')

    s=new('Trois ensembles, trois rôles distincts', '01 / Le problème',
          'On sépare les données avant de commencer à entraîner le réseau.',
          notes=f"5 min. L’ensemble d’apprentissage ({counts['train']} images) est le seul qui produit des gradients. La validation ({counts['val']}) sert à surveiller l’entraînement et à choisir les réglages. Le test ({counts['test']}) est réservé au bilan final. Les légères différences par rapport à 60/20/20 viennent des arrondis de répartition. Le split est stratifié : les proportions des classes sont préservées autant que possible. Si l’on regarde souvent le test pour changer un réglage, il n’est plus une évaluation indépendante.")
    widths=[7.1,2.36,2.36]; xx=.68
    for w,col,cap in zip(widths,[TEAL,ORANGE,INK],['60 %  /  TRAIN','20 %  /  VALIDATION','20 %  /  TEST']):
        rect(s,xx,2.6,w,.7,col,radius=False)
        text(s,xx+.12,2.8,w-.2,.3,cap,13,WHITE,True)
        xx+=w+.1
    for i,(key,title,body) in enumerate([('train','Apprendre','Mettre les poids à jour.'),('val','Choisir','Comparer les réglages.'),('test','Évaluer','Mesurer une fois à la fin.')]):
        box(s,.68+i*4.08,3.65,3.82,2.28,f'{counts[key]} images',title+'\n'+body,[TEAL,ORANGE,INK][i],18)
    takeaway(s,'Même graine, mêmes ensembles. Aucun exemple ne doit apparaître dans deux ensembles.')

    s=new('De l’image au vecteur de nombres', '02 / Le réseau',
          'Une transformation simple, identique pendant l’entraînement et la prédiction.',
          notes="4 min. Aplatir signifie ranger les lignes de l’image à la suite : 8×8 devient 64. L’ordre des pixels doit toujours rester identique. Diviser par 16 ramène les intensités dans [0,1] et aide l’optimisation. Ce facteur est connu à l’avance et n’est pas estimé sur les données ; ce n’est donc pas une fuite de test. Pour une normalisation qui apprend des moyennes et écarts-types, les estimer uniquement sur le train.")
    picture(s,ASSETS/'pixels.png',.7,2.36,4.5,3.95)
    code(s,5.65,2.7,6.98,2.0,'x = digits.data.astype(np.float64) / 16.0\ny = digits.target.astype(np.int64)',15,'Préparation des entrées')
    text(s,5.85,5.05,6.4,1.0,'Une image : (64,)\nUn lot de 32 images : (32, 64)',24,INK,True)

    s=new('Un neurone : pondérer, additionner, activer', '02 / Le réseau',
          'Chaque poids règle l’influence d’une entrée ; le biais décale le résultat.',
          notes="5 min. Faire calculer z à voix haute : 2×0,3 + 1×(−0,4) + 0,1 = 0,3. ReLU renvoie max(0,z), donc 0,3. Avec z=−0,8 elle renvoie 0. Le biais est appris comme les poids. Les couches cachées de notre modèle appliquent toutes ce même calcul, mais chacune avec ses propres poids. Expliquer que la non-linéarité est indispensable : deux transformations affines sans activation se réduisent à une seule transformation affine.")
    box(s,.7,2.6,4.05,3.3,'Somme pondérée','z = w₁x₁ + w₂x₂ + b\n\nExemple :\n2 × 0,3 − 0,4 + 0,1 = 0,3',TEAL,19)
    text(s,4.94,3.68,.8,.8,'→',38,TEAL,True)
    box(s,5.95,2.6,6.65,3.3,'Activation ReLU','a = max(0, z)\n\nz = 0,3   →   a = 0,3\nz = −0,8  →   a = 0',ORANGE,22)
    takeaway(s,'ReLU ajoute une non-linéarité : le réseau peut représenter des relations plus complexes.')

    s=new('Notre réseau : 64 → 32 → 10', '02 / Le réseau',
          'Une couche cachée suffit pour ce premier projet pédagogique.',dark=True,
          notes="4 min. Les 64 entrées sont les pixels. Chaque neurone caché reçoit les 64 entrées. Les 10 sorties correspondent aux chiffres 0 à 9. Le modèle n’a pas 64 neurones entraînables dans une couche d’entrée : les entrées sont simplement les données. Compter les paramètres : 64×32 poids + 32 biais puis 32×10 poids + 10 biais, soit 2 410. Les cercles montrés ici sont représentatifs et non un dessin de tous les neurones.")
    neuron_diagram(s,4.3,2.5,1,1.5)
    for x,n,cap in [(3.53,'64','PIXELS'),(6.38,'32','RELU'),(9.23,'10','SOFTMAX')]:
        text(s,x,5.74,2.0,.68,n,32,MINT,True,align=PP_ALIGN.CENTER)
        text(s,x-.15,6.48,2.3,.32,cap,12,'BCD4CE',True,align=PP_ALIGN.CENTER)
    text(s,.7,2.35,2.55,1.7,'2 410\nparamètres\nà apprendre',25,WHITE,True)
    text(s,.73,6.87,10.6,.16,'Schéma simplifié : tous les neurones ne sont pas représentés.',8,'BCD4CE')

    s=new('Calculer une prédiction, par lots', '02 / Le réseau',
          'NumPy fait les multiplications de matrices avec l’opérateur @.',
          notes="5 min. Lire les dimensions de gauche à droite. B désigne la taille du lot et peut être plus petit pour le dernier lot. X a B lignes et 64 colonnes ; W1 a 64 lignes et 32 colonnes. X@W1 a donc la forme (B,32). Le biais b1 de forme (32,) est ajouté à chaque ligne par broadcasting. Même logique pour la seconde couche. Le résultat final contient B distributions, une par image.")
    for i,(name,formula,shape) in enumerate([('Entrées','X','B × 64'),('Couche cachée','ReLU(X @ W1 + b1)','B × 32'),('Sorties','softmax(H @ W2 + b2)','B × 10')]):
        box(s,.68+i*4.08,2.62,3.81,2.84,name,formula+'\n\n'+shape,TEAL if i!=1 else ORANGE,18)
    rect(s,.7,5.83,11.9,.82,INK)
    text(s,.97,6.03,11.35,.43,'W1.shape == (64, 32)     W2.shape == (32, 10)',19,WHITE,font=CODE)

    s=new('Softmax et erreur : deux rôles différents', '02 / Le réseau',
          'Une distribution pour prédire ; une fonction de coût pour apprendre.',
          notes="5 min. Softmax transforme les scores bruts, appelés logits, en valeurs positives dont la somme vaut 1. Le chiffre choisi est l’argmax. Pour une vraie classe 7, une probabilité 0,8 donne une perte −ln(0,8)≈0,223 ; une probabilité 0,1 donne ≈2,303. La perte est moyennée sur le lot. Le code utilise des log-probabilités stables et soustrait le plus grand logit avant l’exponentielle. Une sortie softmax de 0,9 n’est pas une garantie de 90 % de fiabilité : la calibration n’a pas été évaluée.")
    box(s,.7,2.62,5.76,3.32,'Softmax','pₖ = exp(zₖ) / Σ exp(zⱼ)\n\nLes 10 probabilités totalisent 1.\nLa prédiction choisit le maximum.',TEAL,20)
    box(s,6.68,2.62,5.95,3.32,'Entropie croisée','Pour le vrai chiffre y : −ln(pᵧ)\n\npᵧ = 0,8  →  erreur ≈ 0,223\npᵧ = 0,1  →  erreur ≈ 2,303',ORANGE,20)
    takeaway(s,'Le but est de réduire l’erreur ; le pourcentage de bonnes réponses mesure autre chose.')

    s=new('Apprendre : faire un petit pas dans le bon sens', '02 / Le réseau',
          'Le gradient indique comment une variation des poids change l’erreur.',
          notes="5 min. Présenter w ← w − η×gradient sans dérivation. Si le gradient est positif, on diminue le poids. S’il est négatif, on l’augmente. η est le taux d’apprentissage. Avec w=0,5, η=0,1 et gradient=0,2, le nouveau poids vaut 0,48. Le gradient ne fournit pas une solution magique en un pas ; des mises à jour trop grandes peuvent augmenter l’erreur. Refaire le calcul avec un gradient négatif.")
    rect(s,.72,2.65,11.9,1.35,INK)
    text(s,1.02,2.94,11.1,.7,'poids ← poids − taux × gradient',32,WHITE,True,align=PP_ALIGN.CENTER)
    box(s,.7,4.35,3.8,1.65,'Le poids','Un nombre appris.',body_size=18)
    box(s,4.77,4.35,3.8,1.65,'Le taux η','La taille de chaque pas.',ORANGE,18)
    box(s,8.84,4.35,3.8,1.65,'Le gradient','La direction locale.',body_size=18)
    takeaway(s,'Exemple : 0,5 − 0,1 × 0,2 = 0,48. Le même principe s’applique à tous les paramètres.')

    s=new('La rétropropagation, vue de près', '02 / Le réseau',
          'Option mathématique : propager la sensibilité de l’erreur vers l’entrée.',
          notes="Lecture facultative, 5 min. P est la matrice des probabilités et Y la matrice one-hot des vraies classes. La combinaison softmax/entropie croisée donne dZ2=(P−Y)/B. dW2=HᵀdZ2 ; db2 est la somme des lignes. On remonte ensuite à travers les poids W2, puis on multiplie par le masque z1>0 pour la dérivée de ReLU. À zéro, le code choisit la dérivée nulle. Les gradients doivent tous être calculés avant la moindre mise à jour des poids. Un test de différences finies valide le calcul dans le code livré.")
    code(s,.7,2.5,7.65,3.6,'dz2 = np.exp(log_probs)\ndz2[np.arange(len(y)), y] -= 1\ndz2 /= len(y)\ndw2 = hidden.T @ dz2\ndz1 = (dz2 @ self.w2.T) * (z1 > 0)\ndw1 = x.T @ dz1',15,'Network.loss_and_gradients — extrait')
    text(s,8.7,2.78,3.65,2.8,'1. Erreur en sortie\n\n2. Gradient de W2\n\n3. Traversée de ReLU\n\n4. Gradient de W1',19)
    takeaway(s,'Les biais ont aussi un gradient : somme des gradients du lot, colonne par colonne.')

    s=new('La boucle d’entraînement', '03 / Le code',
          'Un lot = quelques exemples. Une époque = un passage sur tout le train.',
          notes="4 min. Un lot contient ici au plus 32 exemples. Avec 1 078 exemples d’apprentissage, il y a 34 lots par époque, dont un dernier plus petit. À chaque époque, mélanger l’ordre des indices évite un ordre fixe des classes. Le code fait descente de gradient stochastique par mini-lots, pas Adam. Les métriques de train et validation sont recalculées à la fin de chaque époque, sans mise à jour pendant l’évaluation. Le test reste fermé.")
    steps=[('Mélanger','Changer l’ordre\ndes exemples.'),('Prédire','Calculer la sortie\npour un lot.'),('Corriger','Calculer les gradients\net ajuster les poids.'),('Mesurer','Suivre train + validation\naprès chaque époque.')]
    for i,(title,body) in enumerate(steps):
        box(s,.68+i*3.06,2.73,2.84,2.9,f'{i+1}. {title}',body,TEAL if i%2==0 else ORANGE,17)
    takeaway(s,'Réglages du projet : 60 époques · lots de 32 · taux 0,1 · graine 42.')

    s=new('Installer et lancer le projet', '03 / Le code',
          'Exemple Windows / PowerShell, depuis le dossier code du projet décompressé.',
          notes="Manipulation, 8 min. Utiliser Python 3.12, la version testée pour le cours. Ouvrir PowerShell dans le dossier code du ZIP, créer l’environnement et installer requirements.txt. Les commandes utilisent directement le Python de l’environnement : pas besoin de modifier la politique d’activation PowerShell. Le téléchargement des paquets nécessite Internet une fois ; les données digits sont ensuite locales. Les commandes Linux/macOS sont dans le guide. Si la machine n’a pas d’affichage, seuls les tests hors écran peuvent fonctionner.")
    code(s,.7,2.53,11.92,2.64,'py -3.12 -m venv .venv\n.venv\\Scripts\\python.exe -m pip install -r requirements.txt\n.venv\\Scripts\\python.exe app.py',17,'Trois commandes')
    text(s,.9,5.58,11.4,.72,'NumPy : calculs   •   scikit-learn : données et séparation   •   PySide6 : interface',18)
    source(s,'Versions du projet : NumPy 2.2.6 · scikit-learn 1.7.2 · PySide6 6.8.3. Installation desktop requise pour afficher la fenêtre.')

    s=new('Séparer le modèle de son interface', '03 / Le code',
          'On peut entraîner le réseau sans ouvrir de fenêtre.',
          notes="3 min. Ouvrir model.py et localiser load_data, Network.forward, loss_and_gradients, train_epoch et evaluate. Ouvrir app.py et montrer que les boutons appellent le modèle sans réécrire les calculs. Cette séparation permet les tests et les expériences en ligne de commande. test_model.py vérifie la logique, tandis que les exigences de l’interface sont testées hors écran. Ne pas recopier toutes les classes dans un seul gros gestionnaire de bouton.")
    box(s,.7,2.65,3.8,3.2,'model.py','Les données\nLe réseau NumPy\nL’apprentissage\nLa sauvegarde',TEAL,20)
    box(s,4.77,2.65,3.8,3.2,'app.py','La fenêtre Qt\nLes contrôles\nLe travail en arrière-plan\nL’affichage des résultats',ORANGE,18)
    box(s,8.84,2.65,3.8,3.2,'test_model.py','Vérifier les gradients\nMesurer les progrès\nTester la sauvegarde\nContrôler les ensembles',TEAL,18)
    takeaway(s,'Cette architecture facilite la compréhension, les tests et les futures améliorations.')

    s=new('Préparer les ensembles dans Python', '03 / Le code',
          'stratify conserve les proportions des chiffres ; random_state fixe le tirage.',
          notes="Manipulation, 5 min. Faire afficher les shapes de chaque ensemble. Le tableau train contient des indices, pas des images. Le premier split donne 60 % au train et 40 % au reste. Le second partage ce reste en deux, avec stratify=y[rest]. Les mêmes indices permettent de prendre les pixels et les labels correspondants. La graine 43 du second split est seed+1 dans le fichier. Insister sur l’alignement x/y : les mélanger séparément briserait les labels.")
    code(s,.7,2.43,11.92,3.8,'train, rest = train_test_split(\n    np.arange(len(y)), test_size=0.4,\n    stratify=y, random_state=42\n)\nval, test = train_test_split(\n    rest, test_size=0.5,\n    stratify=y[rest], random_state=43\n)',17,'load_data — extrait avec graine explicite')
    takeaway(s,f'Résultat mesuré : train {counts["train"]} · validation {counts["val"]} · test {counts["test"]}.')

    s=new('La propagation avant tient en sept lignes', '03 / Le code',
          'Voici le cœur de Network.forward(), utilisé pour apprendre et pour prédire.',
          notes="Manipulation, 6 min. Associer chaque ligne au schéma 64→32→10. z1 contient les sommes avant ReLU ; hidden les activations. shifted soustrait le maximum par image pour éviter des exponentielles trop grandes. log_probs contient le logarithme des probabilités, donc exp(log_probs) permet de les récupérer. L’argmax des log-probabilités donne la même classe que celui des probabilités, car log est croissant. Les trois tableaux sont rendus pour servir au calcul des gradients.")
    code(s,.7,2.48,11.92,3.61,'z1 = x @ self.w1 + self.b1\nhidden = np.maximum(0, z1)\nlogits = hidden @ self.w2 + self.b2\nshifted = logits - logits.max(axis=1, keepdims=True)\nlog_probs = shifted - np.log(\n    np.exp(shifted).sum(axis=1, keepdims=True))\nreturn z1, hidden, log_probs',16,'Extrait exact de model.py')
    takeaway(s,'La stabilité numérique fait partie de l’algorithme : elle évite de calculer exp(1000).')

    s=new('Entraîner, puis regarder la validation', '03 / Le code',
          'Les réglages sont explicites pour pouvoir refaire la même expérience.',
          notes="Manipulation, 7 min. Le snippet suppose les imports de model.py et NumPy. Le modèle et le générateur aléatoire ont la même graine que la référence. Un appel à train_epoch réalise tous les lots du train. evaluate ne modifie pas les paramètres. Pour une comparaison de taux d’apprentissage, recréer un réseau et un rng pour chaque expérience afin de garder le même point de départ. Ne pas utiliser test à l’intérieur de cette boucle.")
    code(s,.7,2.43,8.57,4.18,'data = load_data(seed=42)\nmodel = Network(seed=42)\nrng = np.random.default_rng(42)\nfor epoch in range(60):\n    model.train_epoch(\n        data["x_train"], data["y_train"],\n        lr=0.1, batch_size=32, rng=rng)\n    val = model.evaluate(\n        data["x_val"], data["y_val"])\n    print(epoch + 1, val)',14.5,'Boucle minimale compatible avec model.py')
    text(s,9.61,2.8,2.81,3.3,'Train :\nmodifie les poids.\n\nValidation :\nmesure seulement.\n\nTest :\nattend la fin.',19)

    s=new('Lire les courbes d’apprentissage', '04 / Observer',
          'Résultats réellement obtenus avec le code fourni, sur une exécution reproductible.',
          notes=f"5 min. Montrer les deux métriques : perte et exactitude. La première mesure la qualité des probabilités, la seconde la proportion de classes correctes. Une même exactitude peut cacher des probabilités différentes et donc une perte différente. Les courbes viennent du fichier support/metrics.json, avec 60 époques, taux 0,1, lots de 32 et graine 42. À la dernière époque, train={pct(last['train_accuracy'])}, validation={pct(last['val_accuracy'])}. Les petites oscillations de validation sont normales ; ne pas promettre une amélioration à chaque époque.")
    picture(s,ASSETS/'loss.png',.67,2.55,6.03,3.65)
    picture(s,ASSETS/'accuracy.png',6.75,2.55,5.96,3.65)
    takeaway(s,f'À 60 époques : train {pct(last["train_accuracy"])}  /  validation {pct(last["val_accuracy"])}.')
    source(s,'Source : expérience locale mesurée ; paramètres complets et 60 points par courbe dans support/metrics.json.')

    s=new('Une bonne courbe peut cacher un problème', '04 / Observer',
          'Le score d’apprentissage seul ne dit pas si le réseau généralise.',
          notes="4 min. Sous-apprentissage : train et validation restent faibles, modèle ou entraînement insuffisant. Surapprentissage : le train progresse tandis que la validation stagne ou se dégrade. Fuite : des exemples de validation/test participent à l’apprentissage ou à un prétraitement appris ; le score devient trompeur. Ne pas conclure à partir d’une seule oscillation. Le taux et le nombre d’époques sont choisis sur validation, pas sur test. L’arrêt anticipé pourrait garder les meilleurs poids de validation ; il n’est pas implémenté ici.")
    for i,(title,body) in enumerate([('Pas assez appris','Train faible\nValidation faible\n\nVérifier les données,\nle taux et les époques.'),('Trop mémorisé','Train élevé\nValidation dégradée\n\nArrêter plus tôt ou\nrégulariser le modèle.'),('Évaluation faussée','Des données se croisent\nou le test guide les choix.\n\nRepartir d’une séparation\ncorrecte des rôles.')]):
        box(s,.7+i*4.07,2.56,3.8,3.68,title,body,TEAL if i==0 else ORANGE,17)
    takeaway(s,'Le bon réflexe : conserver une validation indépendante et garder le test pour le bilan final.')

    s=new('Le bilan final sur des images réservées', '04 / Observer',
          'La matrice de confusion permet de voir quels chiffres sont confondus.',
          notes=f"5 min. Le score final mesuré est {pct(report['test']['accuracy'])} sur {counts['test']} images réservées. Les lignes de la matrice sont les vraies classes, les colonnes les prédictions. Les nombres de la diagonale sont corrects ; les autres indiquent les erreurs. Cette unique répartition n’est pas une garantie universelle, ni un intervalle de confiance. Un classifieur qui choisit uniformément un chiffre au hasard a une espérance de 10 %, mais le réseau aléatoirement initialisé n’a pas nécessairement un score exactement égal à 10 %.")
    picture(s,ASSETS/'confusion.png',.76,2.22,6.55,4.64)
    text(s,8.15,2.8,4.2,1.0,pct(report['test']['accuracy']),43,TEAL,True)
    text(s,8.15,3.94,4.18,1.0,f'de bonnes réponses\nsur {counts["test"]} images de test.',21)
    text(s,8.15,5.38,4.14,.83,'Diagonale = correct\nHors diagonale = confusion',16,MUTED)
    source(s,'Mesure locale : 64→32→10 · SGD · 60 époques · taux 0,1 · lots 32 · graine 42. Score non garanti sur d’autres données.')

    s=new('Une petite interface, quatre actions', '05 / L’interface',
          'L’écran rend l’apprentissage observable sans masquer le fonctionnement du modèle.',
          notes="Démonstration puis manipulation, 6 min. Demander à l’apprenant de lancer un entraînement, observer la courbe, attendre la fin puis effectuer une prédiction. La sélection d’un exemple ne réentraîne pas le réseau. L’évaluation de test est une action explicite, à réserver au bilan après les choix de réglages. Sauvegarder les poids permet de refaire une prédiction plus tard sans refaire l’entraînement. Les chiffres montrés viennent du même jeu 8×8 ; il n’y a pas de reconnaissance de photos arbitraires.")
    picture(s,ASSETS/'interface.png',.68,2.34,8.8,4.35)
    text(s,9.72,2.73,2.81,3.94,'1  Choisir les réglages\n\n2  Lancer l’entraînement\n\n3  Prédire un exemple\n\n4  Sauvegarder le modèle',18,INK,True)

    s=new('Garder la fenêtre réactive', '05 / L’interface',
          'L’entraînement tourne dans un worker ; l’interface reçoit des signaux Qt.',
          notes="5 min. L’événementiel Qt doit continuer à traiter les clics et le dessin de la fenêtre. Une grosse boucle directement dans le callback d’un bouton bloque cette boucle d’événements. Le projet lance le calcul dans un QThread et envoie des métriques à la fenêtre via des signaux. Seul le thread principal modifie les widgets. Une demande d’annulation est vérifiée entre les lots. Le modèle en cours d’entraînement n’est pas utilisé simultanément pour prédire ; les contrôles concernés sont désactivés.")
    box(s,.7,2.67,4.78,2.8,'Thread de l’interface','Boutons et contrôles\nCourbes et résultats\nRéception des signaux',TEAL,21)
    box(s,7.1,2.67,5.52,2.8,'Worker d’entraînement','Boucle sur les époques\nGradients et mises à jour\nÉmission des métriques',ORANGE,21)
    text(s,5.68,3.23,1.2,1.5,'→\n←',35,TEAL,True,align=PP_ALIGN.CENTER)
    takeaway(s,'Le worker calcule ; les slots de la fenêtre mettent les widgets à jour.')
    source(s,'Source : Qt for Python, QThread et mécanisme de signaux/slots.', 'https://doc.qt.io/qtforpython-6/PySide6/QtCore/QThread.html')

    s=new('Relier les calculs aux actions utilisateur', '05 / L’interface',
          'Une action a une responsabilité précise : entraîner, annuler ou prédire.',
          notes="4 min. Le tableau est un contrat de fonctionnement, pas un second code à recopier. app.py contient l’implémentation complète. Expliquer le rôle des signaux progress/result/error, en suivant le bouton d’entraînement jusqu’au worker. Pour une extension, ajouter d’abord la fonction côté modèle, puis l’action côté interface, puis un test. La présence d’un thread ne rend pas automatiquement tout accès concurrent sûr.")
    rows=[('Entraîner','Créer un modèle et démarrer le worker.'),('Progression','Afficher la perte et l’exactitude de validation.'),('Annuler','Demander l’arrêt, puis attendre la fin du worker.'),('Prédire','Calculer predict_proba(image) puis argmax.'),('Évaluer le test','Calculer une mesure finale, sans gradient.')]
    for i,(a,b) in enumerate(rows):
        rect(s,.7,2.47+i*.78,11.9,.64,WHITE if i%2==0 else 'E8F0EB')
        text(s,.94,2.61+i*.78,2.65,.4,a,17,TEAL,True)
        text(s,3.85,2.61+i*.78,8.42,.4,b,17)

    s=new('Sauvegarder ce que le réseau a appris', '05 / L’interface',
          'Un modèle entraîné est principalement un ensemble de poids et de biais.',
          notes="4 min. Le fichier .npz enregistre w1, b1, w2, b2 et une version. Le chargement vérifie les noms, dimensions et valeurs finies ; allow_pickle=False empêche de charger des objets Python sérialisés par pickle. Le fichier ne contient pas l’historique d’entraînement ni le générateur aléatoire : recharger les poids permet l’inférence, pas la reprise exacte de toute la session. L’architecture 64→32→10 et la normalisation /16 doivent être identiques.")
    code(s,.7,2.66,7.1,2.15,'model.save("mon_modele.npz")\nmodel = Network.load("mon_modele.npz")\nprobas = model.predict_proba(image)',16,'API du modèle fourni')
    box(s,8.11,2.66,4.51,3.2,'À conserver ensemble','Les poids\nL’architecture\nLa normalisation\nL’ordre des classes',TEAL,21)
    takeaway(s,'Prédire ne modifie pas les poids. Le modèle ne continue pas à apprendre tout seul.')

    s=new('À vous : trois expériences courtes', '06 / Pratiquer',
          'Notez votre hypothèse avant de lancer le programme. Gardez le test fermé.',
          notes="Exercices, 12 min. A : calculer le nombre de paramètres d’une architecture 64→16→10, sans modifier le code fourni ; réponse sur la diapo de correction. B : sur 20 époques, comparer les taux 0,01 et 0,1 depuis le même seed et noter train/validation. Il n’y a pas de score universel promis ; regarder les mesures. C : fournir le cas fictif train 99,8 % et validation 86 % qui baisse, demander un diagnostic et une action. Les scores de C sont un scénario d’exercice, pas les résultats réels du projet.")
    box(s,.7,2.6,3.8,3.95,'A  Compter','Imaginez 16 neurones cachés.\n\nCombien de poids\net de biais ?',TEAL,17)
    box(s,4.77,2.6,3.8,3.95,'B  Comparer','20 époques, seed 42.\nTaux : 0,01 puis 0,1.\n\nQuel apprentissage\nsemble le plus efficace ?',ORANGE,17)
    box(s,8.84,2.6,3.8,3.95,'C  Diagnostiquer','Scénario fictif :\ntrain = 99,8 %\nvalidation = 86 % et baisse.\n\nQue changer en premier ?',TEAL,17)

    s=new('Exercice code : retrouver les trois choix', '06 / Pratiquer',
          'Recréez le classement affiché par l’interface, puis comparez avec predict_sample().',
          notes="Exercice, 8 min. L’affichage des trois choix est déjà présent dans l’application corrigée fournie : le reconstruire dans la console avant de lire predict_sample. Utiliser une seule image de forme (1,64). predict_proba renvoie (1,10). Prendre la première ligne puis trier les indices. Les apprenants doivent afficher trois classes et leurs scores. Ne pas modifier les labels d’entraînement et ne pas mettre à jour les poids. Ce classement montre les alternatives du réseau, pas une garantie de fiabilité. La correction minimale est donnée sur la prochaine diapositive.")
    code(s,.7,2.63,7.18,2.47,'probas = model.predict_proba(image)[0]\n# 1. Trouver les 3 plus grands scores.\n# 2. Retrouver leurs chiffres.\n# 3. Formater les résultats en pourcentages.',15,'Point de départ')
    box(s,8.2,2.63,4.42,3.24,'Critère de réussite','Trois chiffres distincts,\ntriés par score décroissant.\n\nLes poids ne changent pas.',TEAL,20)
    takeaway(s,'Indice : np.argsort renvoie les indices du plus petit score au plus grand.')

    s=new('Corrections : expliquer le résultat', '06 / Pratiquer',
          'Le raisonnement compte autant que le score obtenu.',
          notes="Correction, 6 min. A : 64×16+16 + 16×10+10 = 1 210. B : le taux 0,1 peut converger plus rapidement ici, mais constater les courbes produites plutôt que l’affirmer sans mesure. Un taux plus élevé peut aussi déstabiliser. C : un écart et une validation qui se dégrade suggèrent un surapprentissage ; garder le meilleur état de validation et envisager régularisation ou réduction de capacité. Le test n’aide pas à choisir. D : le snippet trie les scores puis prend les trois derniers à l’envers. Le choix des 16 neurones n’est pas paramétrable dans l’interface livrée : A est un exercice de calcul.")
    text(s,.75,2.44,5.72,2.38,'A  64 × 16 + 16 + 16 × 10 + 10\n     = 1 210 paramètres.\n\nB  Comparer les pertes de validation.\n\nC  Surapprentissage probable :\n     surveiller la validation et arrêter plus tôt.',18)
    code(s,6.82,2.48,5.78,2.73,'top3 = np.argsort(probas)[-3:][::-1]\nfor chiffre in top3:\n    score = 100 * probas[chiffre]\n    print(f"{chiffre} : {score:.1f} %")',15,'Correction de l’exercice D')
    takeaway(s,'Une explication honnête distingue le score mesuré, une hypothèse et une garantie.')

    s=new('Quand cela ne fonctionne pas', '06 / Pratiquer',
          'Commencez par le symptôme, puis vérifiez une seule hypothèse à la fois.',
          notes="3 min. Le projet fournit des tests. Erreur d’import : vérifier que pip et l’application utilisent le même Python de .venv. Fenêtre absente dans le cloud : un serveur sans affichage ne peut montrer une interface desktop ; le mode offscreen permet seulement une validation technique. NaN : vérifier le taux et les valeurs d’entrée ; notre softmax est stable mais ne compense pas des poids non finis. Mauvais score : normalisation, dimensions, labels et ordre des données. Le test de gradient valide une différence finie sur plusieurs paramètres.")
    rows=[('Module introuvable','Installer requirements.txt avec le Python de .venv.'),('Erreur ou valeurs NaN','Baisser le taux ; vérifier les nombres et les dimensions.'),('L’interface se bloque','Déplacer la boucle d’entraînement dans le worker.'),('Les réponses sont mauvaises','Vérifier /16, forme (1,64), ordre des pixels et des classes.'),('Aucune fenêtre dans le cloud','Lancer sur un bureau ; offscreen sert uniquement à tester.')]
    for i,(a,b) in enumerate(rows):
        rect(s,.7,2.5+i*.78,11.9,.64,WHITE if i%2==0 else 'E8F0EB')
        text(s,.92,2.63+i*.78,3.5,.4,a,14,TEAL,True)
        text(s,4.47,2.63+i*.78,7.97,.4,b,15)

    s=new('Ce que vous avez réellement construit', '07 / Continuer',
          'Un classifieur de chiffres, avec une vraie boucle d’apprentissage.',dark=True,
          notes="Bilan, 3 min. Faire reformuler le chemin pixels→réseau→probabilités→erreur→gradients→poids. Demander ce qui distingue train, validation et test. Le modèle n’est pas un assistant conversationnel, il n’apprend pas automatiquement après chaque prédiction et il est spécialisé dans des chiffres 8×8. Pour aller plus loin : régularisation L2, arrêt anticipé, plus de couches ou réseau convolutif. Changer vers MNIST impose une autre taille d’entrée et un nouveau protocole de données ; il ne suffit pas de charger l’ancien .npz.")
    for i,(t,b) in enumerate([('Vous savez expliquer','Les poids, l’erreur\net la mise à jour.'),('Vous savez vérifier','Un gradient, des courbes\net un score indépendant.'),('Vous savez relier','Un modèle numérique\nà une interface réactive.')]):
        text(s,.78+i*4.04,2.8,3.78,.72,t,23,MINT,True)
        text(s,.78+i*4.04,3.88,3.7,1.2,b,22,WHITE)
    text(s,.78,5.7,11.85,.64,'Prochaine étape : ajouter l’arrêt anticipé en gardant les meilleurs poids de validation.',19,'BCD4CE')

    s=new('Ressources et fichiers du cours', '07 / Continuer',
          'Le PowerPoint est éditable ; les notes de chaque diapositive détaillent les explications.',
          notes="Conclure en montrant le contenu du ZIP. code contient l’application et les tests ; support contient les mesures réelles, le modèle de démonstration, la capture et le générateur de présentation. Le PDF reprend les diapositives pour une lecture facile. Les liens de cette page pointent vers les documentations officielles ; ils complètent le cours. Les mesures du support sont propres à une seule expérience et restent distinctes des scénarios fictifs d’exercice.")
    box(s,.7,2.58,5.23,3.8,'Dans le dossier du cours','code/ : application et tests\nsupport/ : figures et mesures\nlivrables/ : PowerPoint et PDF\nLIRE_MOI.md : démarrage rapide',TEAL,18)
    refs=[('NumPy — prise en main','https://numpy.org/doc/stable/user/absolute_beginners.html'),
          ('scikit-learn — load_digits','https://scikit-learn.org/stable/modules/generated/sklearn.datasets.load_digits.html'),
          ('scikit-learn — train_test_split','https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.train_test_split.html'),
          ('Qt for Python — QThread','https://doc.qt.io/qtforpython-6/PySide6/QtCore/QThread.html'),
          ('scikit-learn — réseaux de neurones','https://scikit-learn.org/stable/modules/neural_networks_supervised.html')]
    for i,(caption,url) in enumerate(refs):
        sh=text(s,6.43,2.8+i*.65,6.03,.45,caption+'  ↗',17,TEAL,True)
        sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url

    for index,(s,dark) in enumerate(slides,1):
        text(s,11.4,7.15,1.2,.24,f'{index:02d} / {len(slides):02d}',10,MINT if dark else TEAL,
             True,align=PP_ALIGN.RIGHT)
    P.core_properties.title='Créer une IA qui apprend en Python'
    P.core_properties.subject='Réseau de neurones NumPy et interface PySide6 — cours en français'
    P.core_properties.author='Cours Python'
    P.core_properties.keywords='Python, intelligence artificielle, NumPy, réseau de neurones, PySide6, cours'
    target=OUT/'Cours_IA_Python_Reseau_Neurones.pptx'
    P.save(target)
    print(f'{len(slides)} diapositives enregistrées : {target}')


if __name__=='__main__':
    build(assets())
